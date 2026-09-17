from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
import os

H = os.path.dirname(os.path.abspath(__file__))
BG = RGBColor(0x0E, 0x11, 0x17); PANEL = RGBColor(0x16, 0x1B, 0x22); INK = RGBColor(0xE6, 0xE8, 0xEC); MUTE = RGBColor(0x9A, 0xA3, 0xB2)
ACC = RGBColor(0x7B, 0x93, 0xFF); OK = RGBColor(0x4C, 0xC3, 0x8A); BAD = RGBColor(0xFF, 0x7B, 0x72); WARN = RGBColor(0xE3, 0xB3, 0x4A); PUR = RGBColor(0xB9, 0x9A, 0xF0)
LIVE = "https://princegarg001.github.io/veridian-it-service-agent/"
REPO = "https://github.com/princegarg001/veridian-it-service-agent"

p = Presentation(); p.slide_width = Inches(13.333); p.slide_height = Inches(7.5)


def txt(s, t, x, y, w, h, size=16, color=INK, bold=False):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame; tf.word_wrap = True
    for i, line in enumerate(t.split("\n")):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = para.add_run(); r.text = line; r.font.size = Pt(size); r.font.color.rgb = color; r.font.bold = bold; r.font.name = "Segoe UI"
    return tb


def slide(title, kicker=None, n=None):
    s = p.slides.add_slide(p.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb = BG
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.12), p.slide_height); bar.fill.solid(); bar.fill.fore_color.rgb = ACC; bar.line.fill.background()
    if kicker: txt(s, kicker, 0.6, 0.35, 12, 0.4, 13, ACC, True)
    txt(s, title, 0.6, 0.7, 12.2, 0.9, 30, INK, True)
    if n: txt(s, f"{n} / 10  ·  Veridian IT Service Agent", 9.6, 7.0, 3.5, 0.35, 10, MUTE)
    return s


def bullets(s, items, x, y, w, h, size=15, gap=6):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame; tf.word_wrap = True
    for i, it in enumerate(items):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); para.space_after = Pt(gap)
        head, _, rest = it.partition("|")
        r = para.add_run(); r.text = head; r.font.bold = True; r.font.size = Pt(size); r.font.color.rgb = ACC; r.font.name = "Segoe UI"
        r = para.add_run(); r.text = " " + rest; r.font.size = Pt(size); r.font.color.rgb = INK; r.font.name = "Segoe UI"


def box(s, x, y, w, h, title, body, color=ACC, ts=15, bs=12):
    b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    b.fill.solid(); b.fill.fore_color.rgb = PANEL; b.line.color.rgb = color; b.adjustments[0] = 0.08
    txt(s, title, x + 0.15, y + 0.1, w - 0.3, 0.4, ts, color, True)
    txt(s, body, x + 0.15, y + 0.5, w - 0.3, h - 0.6, bs, INK)


def img(s, name, x, y, h):
    s.shapes.add_picture(os.path.join(H, "img", name), Inches(x), Inches(y), height=Inches(h))


# 1 Title
s = slide("Veridian IT Service Agent", "AIONOS AGENTIC AI FACTORY · ASSIGNMENT 2: INTERNAL SERVICE AGENT")
txt(s, "An internal IT support agent that's grounded, cites its sources, and can be audited. It understands the issue, finds the policy, asks only the follow-ups that matter, resolves or escalates, and logs every step.", 0.6, 1.8, 11.5, 1.4, 20, MUTE)
for i, (t, b, c) in enumerate([("15 / 15", "data-pack requests triaged", OK), ("4", "active queue tickets assessed", ACC), ("2 + 5", "policy conflicts / status corrections caught", WARN), ("0", "invented policies", BAD)]):
    box(s, 0.6 + i * 3.1, 3.6, 2.9, 1.5, t, b, c, 30, 13)
txt(s, "Live demo:  " + LIVE, 0.6, 5.6, 12, 0.4, 16, ACC, True)
txt(s, "GitHub:  " + REPO, 0.6, 6.05, 12, 0.4, 16, ACC, True)
txt(s, "Prince Garg", 0.6, 6.7, 6, 0.4, 14, MUTE)

# 2 Problem
s = slide("The problem and what the brief requires", "PROBLEM & SCOPE", 2)
txt(s, "Veridian Corp IT receives messy, ambiguous requests. Some are risky (phishing, admin access), some conflict with policy (laptop refresh), and some statuses in the queue are wrong. The agent has to act only on the supplied policies.", 0.6, 1.6, 12, 1, 16, MUTE)
reqs = [("Understand the issue", "Hybrid classifier (rules + optional Claude), entity extraction"),
        ("Find policy / resolution", "Policy map + BM25 retrieval + ticket precedents"),
        ("Sensible follow-ups", "Slot filling. Asks only when the answer changes the decision"),
        ("Resolve simple requests", "Self-service guidance with citations (Wi-Fi, VPN renewal)"),
        ("Escalate risky / unclear", "Security-first override, ungrounded → human, conflicts → human"),
        ("Structured ticket", "JSON: priority, risk, owner, SLA, sources, conflicts, next steps"),
        ("Show source used", "Clickable KB-xx / TK-xxxx citations on every answer"),
        ("Audit trail", "Hash-chained log of every stage and actor, exportable")]
for i, (a, b) in enumerate(reqs):
    box(s, 0.6 + (i % 4) * 3.1, 2.75 + (i // 4) * 2.05, 2.95, 1.85, "✓ " + a, b, OK if i % 2 == 0 else ACC, 14, 12)

# 3 Architecture
s = slide("Architecture of the complete solution", "ARCHITECTURE", 3)
layers = [("Presentation", "index.html · app.js: Agent Chat, Triage Console, Tickets, Audit, KB, Architecture tabs", ACC),
          ("Orchestration", "engine.run(): intake → understand → retrieve → reason → guardrails → act + slot-filling loop", PUR),
          ("Understanding", "Weighted intent classifier + entity extractor. Optional Claude (llm.js) returns schema-checked JSON", PUR),
          ("Knowledge & retrieval", "data.js: KB-01…10 + Asset Mgmt Policy + ticket history. Policy map + BM25 + precedent lookup", OK),
          ("Policy & guardrails", "12 deterministic handlers · risk · conflict detection · ungrounded → escalate · status checks", BAD),
          ("Records", "Structured JSON tickets · hash-chained audit trail · localStorage · JSON export", WARN)]
for i, (t, b, c) in enumerate(layers):
    box(s, 0.6, 1.6 + i * 0.9, 7.2, 0.82, t, b, c, 13, 11)
bullets(s, ["Principle:|The LLM proposes, the policy engine decides. Every approval, route and escalation is deterministic, testable, and cites its source.",
            "Zero-dependency:|static web app on GitHub Pages. Works offline in Rules mode, and hybrid mode adds Claude.",
            "Portable:|the engine is a pure function, so it can sit behind Teams/Slack, ServiceNow, or an API with no rewrite.",
            "Safe fallback:|if the LLM fails or disagrees on a security or privileged intent, the rules result wins."], 8.1, 1.6, 4.9, 5, 14, 10)

# 4 Process flow
s = slide("Process flow: from message to action", "PROCESS FLOW", 4)
steps = [("1 Intake", "chat or queue item + requester + date"), ("2 Understand", "intent, confidence, entities"), ("3 Retrieve", "KB policy + BM25 + precedent"),
         ("4 Reason", "policy handler checks slots"), ("5 Guardrails", "risk, conflict, ungrounded"), ("6 Act", "decision + ticket + audit")]
for i, (t, b) in enumerate(steps):
    box(s, 0.6 + i * 2.1, 1.7, 1.9, 1.3, t, b, BAD if i == 4 else ACC, 14, 11)
txt(s, "↺  If a slot that changes the decision is missing (lockout attempts, laptop age, contractor or full-time, asset tag), the agent asks ONE targeted follow-up and loops back. After 2 failed clarifications it hands over to a human.", 0.6, 3.2, 12, 0.8, 14, WARN)
decs = [("RESOLVED", "Self-service with citation\nREQ-02 guest Wi-Fi\nREQ-05 VPN renewal\nREQ-09 archive mail", OK),
        ("IT_TICKET", "IT must act\nREQ-03 manual unlock\nREQ-13 laptop diagnosis", ACC),
        ("ROUTED", "Other owner\nREQ-07 manager → Finance → IT\nREQ-11 contractor VPN form", PUR),
        ("ESCALATED", "Risk / no authority / conflict\nREQ-01, 04, 08, 10, 14", BAD),
        ("NEEDS_INFO", "Targeted question\nREQ-06 printer\nREQ-12 expense\nREQ-15 vague", WARN)]
for i, (t, b, c) in enumerate(decs):
    box(s, 0.6 + i * 2.5, 4.3, 2.35, 2.4, t, b, c, 15, 12)

# 5 Inputs
s = slide("Inputs, sources and assumptions", "INPUTS & ASSUMPTIONS", 5)
box(s, 0.6, 1.6, 4, 5.2, "Inputs used (only these)", "• Assignment brief: Agentic AI Factory\n\n• Data pack Section 1: KB-01 to KB-10 + Asset Management Policy extract (Finance, Q2 2026)\n\n• Section 2: 15 employee requests REQ-01…15 with recorded status\n\n• Section 3: 10 queue tickets TK-1042…1051 (4 active, 6 closed as precedent)\n\n• No external or invented policy", OK, 16, 13)
bullets(s, ["Dates:|simulation week Mon 21 to Fri 25 Sep 2026. SLAs in business days (no holiday calendar given).",
            "Catalog:|the approved software list isn't supplied, so unconfirmed software is treated as non-catalog (safe default).",
            "Policy conflict:|KB-03 (3 yrs) vs Asset Policy (4 yrs + Finance sign-off) is surfaced to humans, not auto-resolved.",
            "No policy = no action:|admin/privileged access has no KB entry, so escalate with TK-1050 precedent.",
            "REQ-11:|requester treated as the contractor's manager, and this is stated in the reply.",
            "Priority model:|P1 security · P2 employee blocked · P3 standard · P4 informational.",
            "Out of scope:|identity verification, real ITSM/email integration. Tickets are simulated JSON records."], 4.9, 1.6, 8.1, 5.4, 14, 8)

# 6 Chat demo
s = slide("Working agent: grounded conversation", "DEMO · AGENT CHAT", 6)
img(s, "chat_req01.png", 0.6, 1.55, 5.5)
bullets(s, ["REQ-01:|3.5-yr dead laptop",
            "Conflict found:|KB-03 says eligible, the Asset Policy says early (needs Finance sign-off)",
            "Doesn't approve:|escalates to IT verification → IT approval → Finance sign-off",
            "Flags edge case:|the 2-week notice rule doesn't fit a laptop that has already failed",
            "Trace panel:|intake, intent, entities, retrieval scores, precedent TK-1043, guardrails, decision, ticket"], 9.6, 1.6, 3.5, 5.5, 13, 8)

# 7 Triage
s = slide("Triage Console: whole queue in one click", "DEMO · TRIAGE & STATUS RECONCILIATION", 7)
img(s, "triage.png", 0.6, 1.5, 5.7)
bullets(s, ["REQ-03:|'reset queued' is wrong. KB-01 requires a manual unlock",
            "REQ-08:|escalated, but the employee is forwarding the phishing email. Contain it",
            "REQ-04:|SLA window Fri 25 to Tue 29 Sep computed",
            "REQ-06 / REQ-12:|missing asset tag / stalled reply",
            "TK-1043:|HOLD fulfilment. 3.2 yrs approved without Finance sign-off",
            "TK-1048:|link to REQ-08 as a possible same campaign"], 7.3, 1.5, 5.8, 5.6, 14, 9)

# 8 Guardrails & audit
s = slide("Guardrails, escalation and audit trail", "SAFETY & GOVERNANCE", 8)
box(s, 0.6, 1.6, 4, 2.6, "REQ-08 · Phishing (P1 Critical)", "Security-first override. Tells the employee to stop forwarding (KB-09 violation), report to security@veridian-corp.example, warn recipients, and correlate with TK-1048.", BAD, 15, 13)
box(s, 4.8, 1.6, 4, 2.6, "REQ-10 · Urgent admin access", "No KB policy, so the agent won't act. Urgency pressure flagged. Asks for business justification and manager approval. Precedent TK-1050 was rejected for lack of justification.", BAD, 15, 13)
box(s, 9.0, 1.6, 3.8, 2.6, "REQ-14 · Tracking extension", "Non-catalog, so IT Security review (3–5 business days). Marked HIGH risk as a monitoring tool with privacy exposure.", WARN, 15, 13)
box(s, 0.6, 4.4, 6, 2.5, "Audit trail", "Every stage is logged: timestamp, session, actor (employee / agent / LLM / guardrail), stage, detail. Each entry is hash-chained to the previous one, so edits are detectable. Export as JSON.", ACC, 15, 13)
box(s, 6.8, 4.4, 6, 2.5, "Structured ticket", "id · requester · category · decision · priority · risk · assignedTo · status · SLA · kbSources · precedents · policyConflicts · flags · nextSteps · extracted slots. Export as JSON.", OK, 15, 13)

# 9 AI tools
s = slide("AI tools used and how", "AI TOOLS", 9)
tools = [("Claude Code (Claude Opus)", "Build partner", "Parsed the docx and PDF data pack into structured data, designed the pipeline and decision model, wrote the engine, UI, and Node test harness, captured screenshots, generated this deck (python-pptx), and wrote the README. Every rule was checked by hand against the data-pack wording.", ACC),
         ("Claude API (Sonnet), optional", "Runtime NLU", "Hybrid mode (⚙ settings): converts messy employee text into schema-validated {intent, confidence, slots}. Allow-listed slots only. It never approves, routes or escalates, and falls back to rules on failure.", PUR),
         ("Deterministic policy engine", "Decision layer", "12 policy handlers, BM25 retrieval, risk and conflict guardrails. Chosen deliberately over LLM decisions so outcomes are consistent, testable and auditable.", OK),
         ("Node.js scripted evaluation", "Verification", "Replayed all 15 requests and multi-turn dialogues (vague → VPN → contractor, laptop age follow-up, phishing with credentials entered) and fixed a misrouting bug where security must outrank password.", WARN)]
for i, (t, r, b, c) in enumerate(tools):
    box(s, 0.6 + (i % 2) * 6.2, 1.6 + (i // 2) * 2.7, 6, 2.5, f"{t}  ·  {r}", b, c, 15, 13)

# 10 Results
s = slide("Results, limitations and next steps", "OUTCOME", 10)
box(s, 0.6, 1.6, 4, 4.6, "Results", "• 15/15 requests handled with a cited decision or a targeted question\n• 5 escalations, all justified (risk, conflict, no policy)\n• 5 status corrections on the existing queue\n• 2 policy conflicts surfaced (REQ-01, TK-1043)\n• 0 invented policies. Ungrounded requests go to humans\n• Runs with one command or the live link", OK, 16, 13)
box(s, 4.8, 1.6, 4, 4.6, "Limitations", "• Keyword classifier is tuned to this domain (the LLM layer covers messy phrasing)\n• Tickets and audit are in browser storage, not ITSM\n• No identity verification or SSO\n• No holiday calendar for SLAs\n• Catalog contents unknown", WARN, 16, 13)
box(s, 9.0, 1.6, 3.8, 4.6, "Next steps", "• ServiceNow/Jira + Teams integration\n• Embedding retrieval + eval set\n• Policy-owner workflow to reconcile KB-03 vs AMP\n• WORM audit store, PII redaction\n• Human-in-the-loop approval inbox", ACC, 16, 13)
txt(s, "Live: " + LIVE, 0.6, 6.4, 12, 0.4, 15, ACC, True)
txt(s, "Code: " + REPO, 0.6, 6.8, 12, 0.4, 15, ACC, True)

out = os.path.join(H, "Veridian_IT_Service_Agent.pptx"); p.save(out); print(out)
